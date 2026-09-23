import os
import random
import tkinter as tk
import uuid
from tkinter import messagebox, filedialog, ttk
from datetime import datetime, timedelta, timezone
from openpyxl import Workbook, load_workbook
import traceback


def calculate_window_hours(start_str, end_str):
    if start_str == "Closed" or end_str == "Closed" or start_str == end_str:
        return 0.0
    try:
        fmt = "%H:%M"
        t_start = datetime.strptime(start_str, fmt)
        t_end = datetime.strptime(end_str, fmt)

        delta_t = t_end - t_start
        delta_hours = delta_t.total_seconds() / 3600.0
        if delta_hours <= 0:
            delta_hours += 24.0

        return delta_hours
    except ValueError:
        return 0.0


def solve_chore_division(people, chores, disutilities, chore_size, chore_type, people_type, people_schedule, chores_link):

    allocation = {p: [] for p in people}
    housekeeper = []

    running_dis = {p: 0.0 for p in people}
    days_list = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    time = "%H:%M"

    group_assignments = {}

    current_schedule = {}
    current_start_time = {}
    current_consecutive_hours = {}

    for p in people:
        current_schedule[p] = {}
        current_start_time[p] = {}
        current_consecutive_hours[p] = {}
        p_sched = people_schedule.get(p, {})
        for day in days_list:
            day_limit = p_sched.get(day, ("Closed", "Closed"))
            current_schedule[p][day] = calculate_window_hours(day_limit[0], day_limit[1])
            current_start_time[p][day] = day_limit[0]
            current_consecutive_hours[p][day] = 0.0

    link_group_totals = {}
    for chore in chores:
        lk = chores_link.get(chore, "").strip()
        if lk:
            link_group_totals[lk] = link_group_totals.get(lk, 0) + float(chore_size[chore])

    chore_density = []
    for chore in chores:
        d = float(disutilities[chore])
        s = float(chore_size[chore])
        density = d / s if s > 0 else 0
        chore_density.append((chore, d, s, density))

    chore_density.sort(key=lambda x: x[3], reverse=True)

    group_assigned_day = {}

    for chore, dis, size, density in chore_density:
        c_type = chore_type.get(chore, "Any")
        link_group = chores_link.get(chore, "").strip()

        eligible_people = []
        for p in people:
            p_type = people_type.get(p, "Any")
            type_matches = (p_type == "Any" or c_type == "Any" or c_type == p_type)
            if type_matches:
                eligible_people.append(p)

        if not eligible_people:
            housekeeper.append((chore, size, dis))
            continue
        if link_group and link_group in group_assignments:
            best_player = group_assignments[link_group]
        else:
            best_player = min(eligible_people, key=lambda p: running_dis[p])
            if link_group:
                group_assignments[link_group] = best_player

        assigned_day = None

        if link_group and link_group in group_assigned_day:
            target_day = group_assigned_day[link_group]
            if current_schedule[best_player][target_day] >= size:
                assigned_day = target_day
        else:
            required_space = link_group_totals[link_group] if link_group else size
            valid_days = [day for day in days_list if current_schedule[best_player][day] >= required_space]

            if valid_days:
                assigned_day = random.choice(valid_days)
                if link_group:
                    group_assigned_day[link_group] = assigned_day

        if assigned_day is None:
            housekeeper.append((chore, size, dis))
            continue

        start_str = current_start_time[best_player][assigned_day]
        try:
            t_start = datetime.strptime(start_str, time)
            t_end = t_start + timedelta(seconds=int(size * 3600))
            end_str = t_end.strftime(time)
        except ValueError:
            start_str, end_str = "??:??", "??:??"

        allocation[best_player].append((chore, size, dis, assigned_day, start_str, end_str))

        current_consecutive_hours[best_player][assigned_day] += size
        current_schedule[best_player][assigned_day] -= size
        running_dis[best_player] += dis
        next_start_str = end_str

        if current_consecutive_hours[best_player][assigned_day] >= 4.0 and current_schedule[best_player][assigned_day] > 1.0:
            try:
                break_start = datetime.strptime(end_str, time)
                break_end = break_start + timedelta(minutes=60)
                b_start_str = break_start.strftime(time)
                b_end_str = break_end.strftime(time)

                allocation[best_player].append(("Breaktime", 1.0, 0.0, assigned_day, b_start_str, b_end_str))
                current_schedule[best_player][assigned_day] -= 1.0
                next_start_str = b_end_str
                current_consecutive_hours[best_player][assigned_day] = 0.0
            except Exception:
                pass

        current_start_time[best_player][assigned_day] = next_start_str

    day_order = {day: idx for idx, day in enumerate(days_list)}
    for p in allocation:
        allocation[p].sort(key=lambda x: (day_order[x[3]], datetime.strptime(x[5], "%H:%M") if "??" not in x[4] else datetime.min))

    return allocation, housekeeper, current_schedule, running_dis


class Chore:
    def __init__(self, root):
        self.root = root
        self.root.title("Chore")
        self.root.geometry("1500x850")
        self.root.configure(bg="#f8f9fa")

        self.type_options = ["Any", "Databases", "Software", "Maths"]
        self.people = []
        self.people_type = {}
        self.people_schedule = {}
        self.chores = []
        self.chores_type = {}
        self.chores_links = {}

        self.size_entries = {}
        self.disutility_entries = {}
        self.type_dropdown = {}
        self.link_entries = {}

        self. time_increment = ["Closed"] + [f"{h:02d}:{m:02d}" for h in range(24) for m in (0, 30)]

        self.setup_ui_layout()
        self.seed_defaults()

    def setup_ui_layout(self):
        header = tk.Frame(self.root, bg="#f8f9fa", height=60)
        header.pack(fill="x", side="top")
        header.pack_propagate(False)
        tk.Label(header, text="Fair Task Distribution", font=("Helvetica", 14, "bold"), fg="#2228dd", bg="white").pack(pady=15)

        self.left_panel = tk.Frame(self.root, bg="#ffffff", bd=1, relief="solid")
        self.left_panel.pack(side="left", fill="both", expand=True, padx=15, pady=15)

        excel_bar = tk.Frame(self.left_panel, bg="#ffffff", bd=1, relief="groove")
        excel_bar.pack(fill="x", padx=15, pady=5)

        tk.Label(excel_bar, text="Excel Data", font=("Helvetica", 10, "bold"), bg="#f8f8fa", fg="#2228dd").pack(side="left", padx=10)

        tk.Button(excel_bar, text="Export Template", bg="#34485e", fg="white", font=("Helvetica", 10, "bold"), bd=0, padx=10, pady=5, command=self.export_template).pack(side="left", padx=5)

        tk.Button(excel_bar, text="Export Current", bg="#34485e", fg="white", font=("Helvetica", 10, "bold"), bd=0,
                  padx=10, pady=5, command=self.export_current).pack(side="left", padx=5)

        tk.Button(excel_bar, text="Upload Excel Sheet", bg="#34485e", fg="white", font=("Helvetica", 10, "bold"), bd=0,
                  padx=10, pady=5, command=self.upload_sheet).pack(side="left", padx=5)

        config_frame = tk.Frame(self.left_panel, bg="#ffffff", bd=1, pady=5)
        config_frame.pack(fill="x", padx=15)

        #List of Types
        t_box = tk.LabelFrame(config_frame, text="Task Types", font=("Helvetica", 10, "bold"), bg="#ffffff", fg="#2228dd", padx=10, pady=10)
        t_box.pack(side="left", fill="both",expand=True, padx=5)

        t_input_frame = tk.Frame(t_box, bg="#ffffff", bd=1, relief="solid")
        t_input_frame.pack(fill="x",expand=True, pady=2)
        self.t_entry = tk.Entry(t_input_frame, font=("Helvetica", 10), width=12)
        self.t_entry.pack(side="left", pady=5)

        tk.Button(t_input_frame, text="Add Type", bg="#34485e", fg="white", font=("Helvetica", 10),
                  command=self.add_type).pack(side="left", padx=2)

        self.t_listbox = tk.Listbox(t_box, height=4, font=("Helvetica", 9))
        self.t_listbox.pack(fill="x", pady=5)

        tk.Button(t_input_frame, text="Delete Type", bg="#34485e", fg="white", font=("Helvetica", 10),
                  command=self.delete_type).pack(side="left", padx=2)

        #list of Players
        p_box = tk.LabelFrame(config_frame, text="Player Roster", font=("Helvetica", 10, "bold"), bg="#ffffff", fg="#2228dd", padx=10, pady=10)
        p_box.pack(side="left", fill="both",expand=True, padx=5)

        p_input_frame = tk.Frame(p_box, bg="#ffffff", bd=1, relief="solid")
        p_input_frame.pack(fill="x",expand=True, pady=2)

        self.p_entry = tk.Entry(p_input_frame, font=("Helvetica", 10), width=12)
        self.p_entry.pack(side="left", pady=2)

        tk.Label(p_input_frame, text="Type:", font=("Helvetica", 9), bg="#ffffff").pack(side="left", padx=5)
        self.p_type_combo = ttk.Combobox(p_input_frame, values=self.type_options, width=8, state="readonly")
        self.p_type_combo.set("Any")
        self.p_type_combo.pack(side="left", padx=5)

        tk.Button(p_input_frame, text="Add Player", bg="#34485e", fg="white", font=("Helvetica", 10),
                  command=self.add_player).pack(side="left", padx=2)
        tk.Button(p_input_frame, text="Delete Selected Player", bg="#34485e", fg="white", font=("Helvetica", 10),
                  command=self.delete_player).pack(side="left", padx=2)

        p_list_container = tk.Frame(p_box, bg="#ffffff", bd=1, relief="solid")
        p_list_container.pack(fill="x", padx=5)

        self.p_listbox = tk.Listbox(p_list_container, height=4, font=("Helvetica", 9), width=30)
        self.p_listbox.pack(fill="x", expand=True, side="left")

        p_btn_side_panel =  tk.Frame(p_list_container, bg="#ffffff")
        p_btn_side_panel.pack(side="right", fill="y", padx=5)

        tk.Button(p_btn_side_panel, text="Edit", bg="#34485e", fg="white", font=("Helvetica", 9), bd=0, padx=4, pady=2,
                  command=self.open_editor).pack(fill="x", pady=2)
        tk.Button(p_btn_side_panel, text="Time", bg="#34485e", fg="white", font=("Helvetica", 9), bd=0, padx=4, pady=2,
                  command=self.open_hours).pack(fill="x", pady=2)

        #List of Chores
        c_box = tk.LabelFrame(config_frame, text="Chore Roster", font=("Helvetica", 10, "bold"), bg="#ffffff", fg="#2228dd", padx=10, pady=10)
        c_box.pack(side="left", fill="both",expand=True, padx=5)

        c_input_frame = tk.Frame(c_box, bg="#ffffff", bd=1, relief="solid")
        c_input_frame.pack(fill="x",expand=True, pady=2)

        self.c_entry = tk.Entry(c_input_frame, font=("Helvetica", 10), width=12)
        self.c_entry.pack(side="left", pady=2)

        tk.Button(c_input_frame, text="Add Chore", bg="#34485e", fg="white", font=("Helvetica", 10),
                  command=self.add_chore).pack(side="left", padx=2)
        tk.Button(c_input_frame, text="Delete Selected Chore", bg="#34485e", fg="white", font=("Helvetica", 10),
                  command=self.delete_chore).pack(side="left")

        self.c_listbox = tk.Listbox(c_box, height=4, font=("Helvetica", 9), width=30)
        (self.c_listbox.pack(fill="x", pady=5))

        spacer_bar = tk.Frame(self.left_panel, height=2, bg="#e0e0e0")
        spacer_bar.pack(fill="x", padx=15, pady=8)

        self.matrix_outer_frame = tk.LabelFrame(self.left_panel, text="Chore Property Sheet", font=("Helvetica", 10, "bold"), bg="#ffffff",
                                                fg="#2228dd", padx=5, pady=5)
        self.matrix_outer_frame.pack(fill="both", expand=True, padx=15, pady=10)

        self.canvas = tk.Canvas(self.matrix_outer_frame, bd=0, highlightthickness=0, bg="#ffffff")
        self.scrollbar = tk.Scrollbar(self.matrix_outer_frame, orient="vertical", command=self.canvas.yview)
        self.scrollable_inner_frame = tk.Frame(self.canvas, bg="#ffffff")
        self.scrollable_inner_frame.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.create_window((0,0), window=self.scrollable_inner_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=self.scrollbar.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        self.scrollbar.pack(side="right", fill="y")

        right_panel = tk.Frame(self.root, bg="#ffffff", width=450, bd=0, relief="solid")
        right_panel.pack(side="right", fill="both", expand=True, padx=15, pady=15)
        right_panel.pack_propagate(False)

        tk.Label(right_panel, text="Distribute Chores Output", font=("Helvetica", 11, "bold"), bg="#1e1e1e",
                 fg="#00ff00", pady=6).pack(side="top", fill="x")
        self.log_text = tk.Text(right_panel, bg="#151515", fg="#d4d4d4", font=("Helvetica", 10), state="disabled", wrap="word")
        self.log_text.pack(fill="both", expand=True)

        self.control_bar = tk.Frame(self.left_panel, bg="#ffffff", pady=10)
        self.control_bar.pack(side="bottom", fill="x", padx=15)
        self.cal_btn = tk.Button(self.control_bar, text="Distribute Chores", bg="#34485e", fg="white", font=("Helvetica", 10), padx=8, bd=0, state="disabled", command=self.run_distribute)
        self.cal_btn.pack(side="left", fill="x", expand=True)



    def export_template(self):
        file_path = filedialog.asksaveasfilename(defaultextension="FairDivision.xlsx",
                                                 filetypes=[("Excel Files", "*.xlsx")],
                                                 title="FairDivision")
        if not file_path:
            return
        wb = Workbook()
        ws_type = wb.active
        ws_type.title = "FairDivision"
        ws_type.append(["FairDivision"])
        for opt in self.type_options:
            ws_type.append([opt])

        ws_player = wb.create_sheet(title="Players")
        ws_player.append(["Name", "Assignment Type", "Mon_Start", "Mon_End", "Tue_Start", "Tue_End", "Wed_Start", "Wed_End", "Thur_Start", "Thur_End", "Fri_Start", "Fri_End", "Sat_Start", "Sat_End", "Sun_Start", "Sun_End"])
        ws_player.append(["Smith", "Software", "08:00", "18:00", "08:00", "18:00", "08:00", "18:00", "08:00", "18:00", "08:00", "18:00", "Closed", "Closed", "Closed", "Closed"])

        ws_chores = wb.create_sheet(title="Chores")
        ws_chores.append(["Chore Name", "Size", "Disutility", "Chore Type", "linked ID"])
        ws_chores.append(["Computing Basics", 2, 20, "Software", "1"])
        ws_chores.append(["Computing Basics, labs", 2, 10, "Software", "1"])

        try:
            wb.save(file_path)
            messagebox.showinfo("FairDivision", f"FairDivision has been successfully saved at:\n{file_path}")
        except Exception as e:
            messagebox.showerror("FairDivision", f"Could not Save File:{e}")

    def export_current(self):
        file_path = filedialog.asksaveasfilename(defaultextension="FairDivision.xlsx",
                                                 filetypes=[("Excel Files", "*.xlsx")],
                                                 title="Workspace FairDivision")
        if not file_path:
            return
        wb = Workbook()
        ws_type = wb.active
        ws_type.title = "FairDivision"
        ws_type.append(["FairDivision"])
        for opt in self.type_options:
            ws_type.append([opt])

        ws_player = wb.create_sheet(title="Players")
        ws_player.append(
            ["Name", "Assignment Type", "Mon_Start", "Mon_End", "Tue_Start", "Tue_End", "Wed_Start", "Wed_End",
             "Thur_Start", "Thur_End", "Fri_Start", "Fri_End", "Sat_Start", "Sat_End", "Sun_Start", "Sun_End"])
        days = ["Mon", "Tue", "Wed", "Thur", "Fri", "Sat", "Sun"]
        for p in self.people:
            row = [p, self.people_type.get(p, "Any")]
            sched = self.people_schedule.get(p, {d: ("Closed", "Closed") for d in days})
            for d in days:
                row.extend([sched.get(d, ("Closed", "Closed"))[0], sched.get(d, ("Closed", "Closed"))[1]])
            ws_player.append(row)

        ws_chores = wb.create_sheet(title="Chores")
        ws_chores.append(["Chore Name", "Size", "Disutility", "Chore Type", "linked ID"])
        for chore in self.chores:
            sz = self.size_entries[chore].get().strip() if chore in self.size_entries else "0"
            ds = self.disutility_entries[chore].get().strip() if chore in self.disutility_entries else "0"
            tp = self.type_dropdown[chore].get().strip() if chore in self.type_dropdown else "Any"
            lk = self.link_entries[chore].get().strip() if chore in self.link_entries else ""
            ws_chores.append([chore, float(sz), float(ds), tp, lk])

        wb.save(file_path)
        messagebox.showinfo("Success", f"FairDivision has been successfully saved at:\n{file_path}")


    def upload_sheet(self):
        file_path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx")], title="Workspace FairDivision")
        if not file_path:
            return
        try:
            wb = load_workbook(file_path, data_only=True)
            if "FairDivision" in wb.sheetnames:
                self.type_options = ["Any"]
                for row in wb["FairDivision"].iter_rows(min_row=2, values_only=True):
                    if row and row[0]:
                        val = str(row[0]).strip()
                        if val and val not in self.type_options:
                            self.type_options.append(val)
                self.update_type_list_view()

            self.people.clear()
            self.people_type.clear()
            self.people_schedule.clear()
            self.p_listbox.delete(0, tk.END)
            self.chores.clear()
            self.c_listbox.delete(0, tk.END)

            days = ["Mon", "Tue", "Wed", "Thur", "Fri", "Sat", "Sun"]
            temp_people = {}
            temp_sizes, temp_disutilities, temp_types, temp_links = {}, {}, {}, {}

            if "Players" in wb.sheetnames:
                for row in wb["Players"].iter_rows(min_row=2, values_only=True):
                    if not row or not row[0]:
                        continue
                    name = str(row[0]).strip()
                    p_type = str(row[1]).strip() if len(row) > 1 and row[1] else "Any"
                    self.people.append(name)
                    self.people_type[name] = p_type
                    self.people_schedule[name] = {}
                    for idx, day in enumerate(days):
                        start, end = 2 + (idx * 2), 3 + (idx * 2)
                        self.people_schedule[name][day] = (str(row[start]).strip() if len(row) > start and row[start] else "Closed",
                                                            str(row[end]).strip() if len(row) > end and row[end] else "Closed")
                    self.p_listbox.insert(tk.END, f"{name} [{p_type}]")

            chores_data = {}
            for row in wb["Chores"].iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                c_name = str(row[0]).strip()
                self.chores.append(c_name)
                self.c_listbox.insert(tk.END, c_name)
                temp_sizes[c_name] = row[1] if row[1] is not None else 0.0
                temp_disutilities[c_name] = row[2] if row[2] is not None else 0.0
                temp_types[c_name] = str(row[3]).strip() if len(row) > 3 and row[3] else "Any"
                temp_links[c_name] = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""

            self.generate_table()
            self.root.update_idletasks()

            for c in self.chores:
                if c in self.size_entries:
                    try:
                        self.size_entries[c].delete(0, tk.END)
                        self.size_entries[c].insert(0, str(temp_sizes.get(c, "0")))
                    except Exception as e:
                        print(f"error updating size of {c}: {e}")

                    try:
                        self.disutility_entries[c].delete(0, tk.END)
                        self.disutility_entries[c].insert(0, str(temp_disutilities.get(c, "0")))
                    except Exception as e:
                        print(f"error updating disutility of {c}: {e}")

                    try:
                        if c in temp_types and temp_types[c] in self.type_options:
                            self.type_dropdown[c].set(temp_types[c])
                        else:
                            self.type_dropdown[c].set("Any")
                    except Exception as e:
                        print(f"error updating type of {c}: {e}")

                    try:
                        self.link_entries[c].delete(0, tk.END)
                        self.link_entries[c].insert(0, str(temp_links.get(c, "0")))
                    except Exception as e:
                        print(f"error updating link of {c}: {e}")
        except Exception as e:
            messagebox.showerror("FairDivision", f"{e}")
            traceback.print_exc()




    def update_type_list_view(self):
        self.t_listbox.delete(0, tk.END)
        for row in self.type_options:
            if row != "Any":
                self.t_listbox.insert(tk.END, f"{row}")
            self.p_type_combo.config(values=self.type_options)


    def generate_table(self):
        saved_data = {}
        if hasattr(self, "size_entries") and self.size_entries:
            for chore in self.chores:
                if chore in self.size_entries:
                    saved_data[chore] = {
                        "size": self.size_entries[chore].get(),
                        "disutility": self.disutility_entries[chore].get(),
                        "type": self.type_dropdown[chore].get(),
                        "link": self.link_entries[chore].get()
                    }

        for widget in self.scrollable_inner_frame.winfo_children():
            widget.destroy()

        if len(self.people) < 1 or len(self.chores) < 1:
            self.cal_btn.config(state="disabled", bg="red")
            return

        self.size_entries, self.disutility_entries, self.type_dropdown, self.link_entries = {}, {}, {}, {}

        tk.Label(self.scrollable_inner_frame, text="Chores Name", font=("Arial", 9, "bold"), width=16, anchor="w").grid(row=0, column=0, padx=5, pady=5)
        tk.Label(self.scrollable_inner_frame, text="Size (Hours)", font=("Arial", 9, "bold"), width=12).grid(row=0, column=1, padx=5, pady=5)
        tk.Label(self.scrollable_inner_frame, text="Disutility", font=("Arial", 9, "bold"), width=14).grid(row=0, column=2, padx=5, pady=5)
        tk.Label(self.scrollable_inner_frame, text="Chore Type", font=("Arial", 9, "bold"), width=14).grid(row=0, column=3,padx=5,pady=5)
        tk.Label(self.scrollable_inner_frame, text="Group ID", font=("Arial", 9, "bold"), width=14).grid(row=0,column=4,padx=5,pady=5)

        for c_idx, chore in enumerate(self.chores):
            tk.Label(self.scrollable_inner_frame, text=chore, font=("Arial", 9, "bold"), fg="#16a085").grid(row=c_idx+1, column=0, padx=5, pady=5, sticky="w")

            se = tk.Entry(self.scrollable_inner_frame, width=8, justify="center")
            if chore in saved_data:
                se.insert(0, saved_data[chore]["size"])
            else:
                se.insert(0, "2")
            se.grid(row=c_idx+1, column=1, padx=5)
            self.size_entries[chore] = se

            de = tk.Entry(self.scrollable_inner_frame, width=8, justify="center")
            if chore in saved_data:
                de.insert(0, saved_data[chore]["disutility"])
            else:
                de.insert(0, "25")
            de.grid(row=c_idx+1, column=2, padx=5)
            self.disutility_entries[chore] = de

            td = ttk.Combobox(self.scrollable_inner_frame, values=self.type_options, width=10, justify="center")
            if chore in saved_data:
                td.insert(0, saved_data[chore]["type"])
            else:
                td.set("Any")
            td.grid(row=c_idx+1, column=3, padx=5)
            self.type_dropdown[chore] = td

            lid = tk.Entry(self.scrollable_inner_frame, width=10, justify="center", bg="#16a085")
            if chore in saved_data:
                lid.insert(0, saved_data[chore]["link"])
            lid.grid(row=c_idx+1, column=4, padx=5)
            self.link_entries[chore] = lid

        self.cal_btn.config(state="normal", bg="#34485e")




    def add_type(self):
        t = self.t_entry.get().strip()
        if t and t not in self.type_options:
            self.type_options.append(t)
            self.update_type_list_view()
            self.t_entry.delete(0, tk.END)
            self.refresh_type_dropdown()

    def delete_type(self):
        idx = self.t_listbox.curselection()
        if not idx:
            return
        txt = self.t_listbox.get(idx[0]).strip()
        if txt in self.type_options:
            self.type_options.remove(txt)
            self.update_type_list_view()
            self.refresh_type_dropdown()

    def refresh_type_dropdown(self):
        for chore in self.chores:
            if chore in self.type_dropdown:
                current_value = self.type_dropdown[chore].get()
                self.type_dropdown[chore]['values'] = self.type_options
                if current_value not in self.type_options:
                    self.type_dropdown[chore].set("Any")

        update = False
        for player in self.people:
            if player in self.people_type:
                current_value = self.people_type[player]
                if current_value not in self.type_options:
                    self.people_type[player] = "Any"
                    update = True
        if update:
            self.update_player_list_display()
        self.p_type_combo.config(values=self.type_options)
        if self.p_type_combo.get() not in self.type_options:
            self.p_type_combo.set("Any")


    def update_player_list_display(self):
        self.p_listbox.delete(0, tk.END)
        for player in self.people:
            p_type = self.people_type.get(player, "Any")
            self.p_listbox.insert(tk.END, f"{player} [{p_type}]")


    def add_player(self):
        p, t = self.p_entry.get().strip(), self.p_type_combo.get()
        if p and p not in self.people:
            self.people.append(p)
            self.people_type[p] = t
            self.people_schedule[p] = {d: ("08:00", "18:00") for d in ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]}
            self.p_listbox.insert(tk.END, f"{p} [{t}]")
            self.p_entry.delete(0, tk.END)
            self.generate_table()

    def delete_player(self):
        idx = self.p_listbox.curselection()
        if not idx:
            return
        n = self.p_listbox.get(idx[0]).split(" [")[0].strip()
        if n in self.people:
            self.people.remove(n)
        self.p_listbox.delete(idx[0])
        self.generate_table()

    def add_chore(self):
        c = self.c_entry.get().strip()
        if c and c not in self.chores:
            self.chores.append(c)
            self.c_listbox.insert(tk.END, c)
            self.c_entry.delete(0, tk.END)
            self.generate_table()

    def delete_chore(self):
        idx = self.c_listbox.curselection()
        if not idx:
            return
        c =self.c_listbox.get(idx[0])
        self.chores.remove(c)
        self.c_listbox.delete(idx[0])
        self.generate_table()

    def open_editor(self):
        idx = self.p_listbox.curselection()
        if not idx:
            return
        old_name = self.p_listbox.get(idx[0]).split(" [")[0].strip()
        popup = tk.Toplevel(self.root, bg="#34485e")
        popup.geometry("300x200")
        popup.title("Edit Profile")
        popup.grab_set()

        tk.Label(popup, text="Name:").pack()
        ne = tk.Entry(popup)
        ne.insert(0, old_name)
        ne.pack()
        tk.Label(popup, text="Type:").pack()
        te = ttk.Combobox(popup, values=self.type_options, state="readonly")
        te.set(self.people_type[old_name])
        te.pack()

        def save():
            new_name = ne.get().strip()
            if new_name:
                self.people[self.people.index(old_name)] = new_name
                self.people_type[new_name] = te.get()
                self.people_schedule[new_name] = self.people_schedule.pop(old_name)
                self.p_listbox.delete(idx[0])
                self.p_listbox.insert(idx[0], f"{new_name} [{te.get()}]")
                popup.destroy()
                self.generate_table()

        tk.Button(popup, text="Save", command=save).pack()

    def open_hours(self):
        idx = self.p_listbox.curselection()
        if not idx:
            return
        name = self.p_listbox.get(idx[0]).split(" [")[0].strip()
        popup = tk.Toplevel(self.root, bg="#34485e")
        popup.geometry("400x450")
        popup.title("Hours")
        popup.grab_set()
        days = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
        tracker = {}
        for day in days:
            f = tk.Frame(popup)
            f.pack(pady=2)
            tk.Label(f, text=f"{day}:", width=6).pack(side="left")
            start = ttk.Combobox(f, values=self.time_increment, width=8, state="readonly")
            start.set(self.people_schedule[name][day][0])
            start.pack(side="left")
            end = ttk.Combobox(f, values=self.time_increment, width=8, state="readonly")
            end.set(self.people_schedule[name][day][1])
            end.pack(side="left")
            tracker[day] = (start, end)

        def save():
            for day in days:
                self.people_schedule[name][day] = (tracker[day][0].get(), tracker[day][1].get())
            popup.destroy()
        tk.Button(popup, text="Save", command=save).pack(pady=10)



    def run_distribute(self):
        disutilities, chore_size = {}, {}
        self.chores_links.clear()

        try:
            for chore in self.chores:
                size = float(self.size_entries[chore].get().strip() or 0)
                if size <= 0:
                    raise ValueError(f"Chore '{chore}' need to be greater than 0")
                chore_size[chore] = size

                dis = float(self.disutility_entries[chore].get().strip() or 0)
                disutilities[chore] = dis

                self.chores_type[chore] = self.type_dropdown[chore].get()
                self.chores_links[chore] = self.link_entries[chore].get().strip()

        except ValueError as err:
            messagebox.showerror("Error", str(err))
            return

        allocation, housekeeper, _, _ = solve_chore_division(self.people, self.chores, disutilities, chore_size, self.chores_type, self.people_type,
                                                             self.people_schedule, self.chores_links)

        self.log_text.config(state="normal")
        self.log_text.delete("1.0", tk.END)
        self.log_text.config(state="disabled")

        for person in self.people:
            self.append_log(f"Assigned to {person}")
            alist = allocation.get(person, [])
            if not alist:
                self.append_log("--No assignments--")
            for chore, size, dis, day, start, end in alist:
                self.append_log(f"Assigned to {chore} - Day: {day} | Shift: {start}-{end} ({size}h)")

        self.open_schedule_popup(allocation, housekeeper)

    def append_log(self, data):
        self.log_text.config(state="normal")
        self.log_text.insert(tk.END, data + "\n")
        self.log_text.see(tk.END)
        self.log_text.config(state="disabled")

    def open_schedule_popup(self, allocation, housekeeper):
        popup = tk.Toplevel(self.root, bg="#34485e")
        popup.title("Schedule")
        popup.geometry("1320x750")
        popup.grab_set()

        header_frame = tk.Frame(popup, bg="#34485e", pady=10)
        header_frame.pack(fill="x")
        tk.Label(header_frame, text="Allocation", font=("Arial", 12, "bold"), bg="#2c3e50", fg="white").pack()
        toolbar_frame = tk.Frame(popup, bg="#34485e", pady=8, bd=1, relief="groove")
        toolbar_frame.pack(fill="x", padx=15, pady=5)

        days_headers = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
        days_map = {"Mon": 0, "Tue": 1, "Wed": 2, "Thu": 3, "Fri": 4, "Sat": 5, "Sun": 6}

        def export_table_to_excel():
            file_path = filedialog.asksaveasfilename(defaultextension="FairDivision.xlsx",
                                                     filetypes=[("Excel Files", "*.xlsx")],
                                                     title="Distribution of FairDivision")
            if not file_path:
                return
            try:
                wb = Workbook()
                ws_out = wb.active
                ws_out.title = "FairDivision Schedule"
                ws_out.append(["Name"] + days_headers)

                for person in self.people:
                    schedule_row_list = {day: [] for day in days_headers}
                    assigned_item = allocation.get(person, [])
                    for item in assigned_item:
                        chore, size, dis, day, start, end = item
                        schedule_row_list[day].append(f"{chore}  ({start}-{end})")

                    row_data = [person]
                    for day in days_headers:
                        row_data.append(", ".join(schedule_row_list[day]) if schedule_row_list[day] else "Free")
                    ws_out.append(row_data)

                if housekeeper:
                    ws_out.append([])
                    ws_out.append(["Unassigned Chores"])
                    for item in housekeeper:
                        ws_out.append([f"{item[0]} ({item[1]} hours)"])
                wb.save(file_path)
                messagebox.showinfo("Success", f"Schedule exported to {file_path}")
            except Exception as err:
                messagebox.showerror("export Error", str(err))

        def write_single_person_ics(person_name, assigned_item):
            valid_chores = [item for item in assigned_item if "Break" not in item[0]]
            if not valid_chores:
                messagebox.showwarning("Empty Schedule", f"{person_name} has no assigned items")
                return

            safe_name = "".join(x for x in person_name if x.isalnum() or x in (' ', '_', '-')).strip()
            save_path = filedialog.asksaveasfilename(
                defaultextension=".ics",
            filetypes=[("ICalender Files", "*.ics")],
            initialfile=f"Timetable_{safe_name}.ics",
            title=f"Timetable for {safe_name}")
            if not save_path:
                return

            today = datetime.now()
            try:
                with open(save_path, "w", encoding="utf-8", newline="\r\n") as ics:
                    ics.write("BEGIN:VCALENDER\r\n")
                    ics.write("VERSION:2.0\r\n")
                    ics.write(f"PRODID:-//Fair Distribution Tasks//{safe_name} Tracker//EN\r\n")
                    ics.write("CALSCALE:GREGORIAN\r\n")
                    ics.write("METHOD:PUBLISH\r\n")

                    for item in valid_chores:
                        chore, size, dis, day, start, end = item

                        target_day_num = days_map.get(day, 0)
                        day_ahead = target_day_num - today.weekday()
                        if day_ahead <= 0: day_ahead +=7
                        target_date = today + timedelta(days=day_ahead)
                        date_str = target_date.strftime("%Y/%m/%d")

                        clean_start = start.replace(":", "") + "00"
                        clean_end = end.replace(":", "") + "00"

                        uid_str = f"{uuid.uuid4()}@{safe_name.lower()}.FairTask.local"
                        timestamp_str = datetime.now(timezone.utc).strftime("%Y/%m/%dT%H:%M:%SZ")

                        ics.write("BEGIN:VEVENT\r\n")
                        ics.write(f"UID:{uid_str}\r\n")
                        ics.write(f"DTSTAMP:{timestamp_str}\r\n")
                        ics.write(f"SUMMARY:{chore}\r\n")
                        ics.write(f"DTSTART;TZID=Local:{date_str}T{clean_start}\r\n")
                        ics.write(f"DTSEND;TZID=Local:{date_str}T{clean_end}\r\n")
                        ics.write(f"DESCRIPTION:Duration: {size} hours\\nDisutility Score: {dis}\r\n")
                        ics.write("SEQUENCE:0\r\n")
                        ics.write("STATUS:CONFIRMED\r\n")
                        ics.write("TRAMSP:OPAQUE\r\n")
                        ics.write("END:VEVENT\r\n")
                    ics.write("END:VCALENDER\r\n")
                messagebox.showinfo("Success", f"Schedule saved calender to {save_path}")
            except Exception as ex:
                messagebox.showerror("export Error", f"Failed to save calender file:\n{str(ex)}")

        def export_all_calender():
            folder_path = filedialog.askdirectory(title="Select Folder")
            if not folder_path:
                return
            today = datetime.now()
            files_created = []

            try:
                for person, assigned_item in allocation.items():
                    valid_chores = [item for item in assigned_item if "Break" not in item[0]]
                    if not valid_chores:
                        continue
                    safe_name = "".join(x for x in person if x.isalnum() or x in (' ', '_', '-')).strip()
                    file_name = f"TimeTable_{safe_name}.ics"
                    full_path = os.path.join(folder_path, file_name)

                    with open(full_path, "w", encoding="utf-8", newline="\r\n") as ics:
                        ics.write("BEGIN:VCALENDER\r\n")
                        ics.write("VERSION:2.0\r\n")
                        ics.write(f"PRODID:-//Fair Distribution Tasks//{safe_name} Tracker//EN\r\n")
                        ics.write("CALSCALE:GREGORIAN\r\n")
                        ics.write("METHOD:PUBLISH\r\n")

                        for item in valid_chores:
                            chore, size, dis, day, start, end = item

                            target_day_num = days_map.get(day, 0)
                            day_ahead = target_day_num - today.weekday()
                            if day_ahead <= 0: day_ahead += 7
                            target_date = today + timedelta(days=day_ahead)
                            date_str = target_date.strftime("%Y/%m/%d")

                            clean_start = start.replace(":", "") + "00"
                            clean_end = end.replace(":", "") + "00"

                            uid_str = f"{uuid.uuid4()}@{safe_name.lower()}.FairTask.local"
                            timestamp_str = datetime.now(timezone.utc).strftime("%Y/%m/%dT%H:%M:%SZ")

                            ics.write("BEGIN:VEVENT\r\n")
                            ics.write(f"UID:{uid_str}\r\n")
                            ics.write(f"DTSTAMP:{timestamp_str}\r\n")
                            ics.write(f"SUMMARY:{chore}\r\n")
                            ics.write(f"DTSTART;TZID=Local:{date_str}T{clean_start}\r\n")
                            ics.write(f"DTSEND;TZID=Local:{date_str}T{clean_end}\r\n")
                            ics.write(f"DESCRIPTION:Duration: {size} hours\\nDisutility Score: {dis}\r\n")
                            ics.write("SEQUENCE:0\r\n")
                            ics.write("STATUS:CONFIRMED\r\n")
                            ics.write("TRAMSP:OPAQUE\r\n")
                            ics.write("END:VEVENT\r\n")
                        ics.write("END:VCALENDER\r\n")
                    files_created.append(file_name)

                messagebox.showinfo("All Calenders Saved", f"{len(files_created)} files created\ntimetable files in:\n{folder_path}\nOne for yeah player with a take.")
            except Exception as ex:
                messagebox.showerror("export Error", f"Failed to save calender file:\n{str(ex)}")

        tk.Button(toolbar_frame, text="Export All to excel", bg="#27ae60", fg="white", font=("Arial", 10, "bold"), bd=0, padx=12, pady=5, command=export_table_to_excel).pack(side="left", padx=10)
        tk.Button(toolbar_frame, text="Save .ics for All Individuals", bg="#27ae60", fg="white", font=("Arial", 10, "bold"), bd=0, padx=12, pady=5, command=export_all_calender).pack(side="left", padx=5)

        grid_outer = tk.Frame(popup, bg="#ffffff", padx=15, pady=5)
        grid_outer.pack(expand=True, fill="both")

        canvas = tk.Canvas(grid_outer, bg="#ffffff", bd=0, highlightthickness=0)
        vsb = tk.Scrollbar(grid_outer, orient="vertical", command=canvas.yview)
        hsb = tk.Scrollbar(grid_outer, orient="horizontal", command=canvas.xview)
        scroll_content = tk.Frame(grid_outer, bg="#eef2f5")

        scroll_content.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0,0), window=scroll_content, anchor="nw")
        canvas.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        canvas.pack(side="top", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        hsb.pack(side="bottom", fill="x")

        all_headers = ["Name"] + days_headers + ["Action"]
        col_widths = [110, 150, 150, 150, 150, 150, 150, 150, 160]

        for idx, header_text in enumerate(all_headers):
            hf = tk.Frame(scroll_content, bg="#3449ee", bd=1, relief="raised", height=32, width=col_widths[idx])
            hf.grid(row=0, column=idx, sticky="nsew")
            hf.grid_propagate(False)
            tk.Label(hf, text=header_text, font=("Arial", 9, "bold"), fg="white", bg="#3449ee").pack(expand=True)

        for r_idx, person in enumerate(self.people, start=1):
            schedule_row_list = {day: [] for day in days_headers}
            assigned_items = allocation.get(person, [])
            for item in assigned_items:
                chore, size, dis, day, start, end = item
                schedule_row_list[day].append((f"{chore}\n ({start}-{end})"))

            cell_f = tk.Frame(scroll_content, bg="#ffffff", bd=1, relief="groove", height=80, width=col_widths[0])
            cell_f.grid(row=r_idx, column=0, sticky="nsew")
            cell_f.grid_propagate(False)
            tk.Label(cell_f, text=person, font=("Arial", 10, "bold"), bg="white", fg="#3449ee").pack(expand=True)

            for d_idx, day in enumerate(days_headers, start=1):
                cell_f = tk.Frame(scroll_content, bg="white", bd=1, relief="groove", height=80, width=col_widths[d_idx])
                cell_f.grid(row=r_idx, column=d_idx, sticky="nsew")
                cell_f.grid_propagate(False)

                txt = "\n-------\n".join(schedule_row_list[day]) if schedule_row_list[day] else "Free"
                lbl_color = "#e31c37" if not schedule_row_list else "#1ce3c8"
                tk.Label(cell_f, text=txt, font=("Arial", 8), bg="white", fg=lbl_color, justify="center").pack(expand=True)

            cell_f = tk.Frame(scroll_content, bg="#fcfcfc", bd=1, relief="groove", height=80, width=col_widths[-1])
            cell_f.grid(row=r_idx, column=len(all_headers) - 1, sticky="nsew")
            cell_f.grid_propagate(False)

            calender_btn = tk.Button(cell_f, text="Download Schedule", font=("Arial", 9, "bold"), bg="#b44b98", fg="white", bd=0, padx=8, pady=6, command=lambda p=person, items=assigned_items: write_single_person_ics(p, items))
            calender_btn.pack(expand=True, padx=5, pady=5)

        if housekeeper:
            keeper_f = tk.Label(popup, text="Housekeeper Overflow:", font=("Arial", 9, "bold"), fg="#e74c3c", bg="white", padx=10, pady=5)
            keeper_f.pack(fill="x", padx=15, pady=(0, 15))
            keeper_txt = "  ".join(f"{item[0]} ({item[1]})" for item in housekeeper)
            tk.Label(keeper_f, text=keeper_txt, font=("Arial", 10), fg="#e74c3c", bg="white", anchor="w").pack(fill="x")


    def seed_defaults(self):
        self.update_type_list_view()
        days = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

        startup_people = [("Andrea", "Software"), ("Ian", "Maths")]
        for p, t in startup_people:
            self.people.append(p)
            self.people_type[p] = t
            self.people_schedule[p] = {
                d: ("08:00", "18:00") if d in ["Mon", "Tue", "Wed", "Thu", "Fri"] else ("Closed", "Closed")
                for d in days
            }
            self.p_listbox.insert(tk.END, f"{p} [{t}]")

        startup_chores = ["Computer Basic's", "Software Fundamentals", "Software Fundamentals Lab", "Advanced Mathematics Theory", "Advanced Mathematics Theory Lab"]
        for c in startup_chores:
            self.chores.append(c)
            self.c_listbox.insert(tk.END, c)



        self.generate_table()
        if "Software Fundamentals" in self.type_dropdown:
            self.type_dropdown["Software Fundamentals"].set("Software")
        if "Software Fundamentals" in self.link_entries:
            self.link_entries["Software Fundamentals"].delete(0, tk.END)
            self.link_entries["Software Fundamentals"].insert(0, "1")

        if "Software Fundamentals Lab" in self.type_dropdown:
            self.type_dropdown["Software Fundamentals Lab"].set("Software")
        if "Software Fundamentals Lab" in self.link_entries:
            self.link_entries["Software Fundamentals Lab"].delete(0, tk.END)
            self.link_entries["Software Fundamentals Lab"].insert(0, "1")

        if "Advanced Mathematics Theory" in self.type_dropdown:
            self.type_dropdown["Advanced Mathematics Theory"].set("Maths")
        if "Advanced Mathematics Theory" in self.link_entries:
            self.link_entries["Advanced Mathematics Theory"].delete(0, tk.END)
            self.link_entries["Advanced Mathematics Theory"].insert(0, "2")

        if "Advanced Mathematics Theory" in self.type_dropdown:
            self.type_dropdown["Advanced Mathematics Theory Lab"].set("Maths")
        if "Advanced Mathematics Theory Lab" in self.link_entries:
            self.link_entries["Advanced Mathematics Theory Lab"].delete(0, tk.END)
            self.link_entries["Advanced Mathematics Theory Lab"].insert(0, "2")


if __name__ == "__main__":
    root = tk.Tk()
    app = Chore(root)
    root.mainloop()